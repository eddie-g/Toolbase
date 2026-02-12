#!/usr/bin/env python3
"""
Convert a PDF to Excel (.xlsx) format.

Usage:
    python3 convert_to_excel.py <input_pdf> <output_xlsx> --mode tables|all [--merge-cells] [--no-merge-cells] [--sheet-per-page] [--single-sheet] [--json]
"""

import sys
import os
import argparse
import json
import traceback
import re

try:
    import fitz  # PyMuPDF
except ImportError:
    print(json.dumps({"success": False, "error": "PyMuPDF is not installed. Run: pip install PyMuPDF"}))
    sys.exit(1)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print(json.dumps({"success": False, "error": "openpyxl is not installed. Run: pip install openpyxl"}))
    sys.exit(1)


def find_tables_on_page(page):
    """
    Detect table structures on a page using line analysis.
    Returns a list of tables, each being a list of rows,
    each row being a list of cell texts.
    """
    try:
        tables = page.find_tables()
        if tables and tables.tables:
            result = []
            for table in tables:
                rows = []
                for row in table.extract():
                    cells = []
                    for cell in row:
                        cells.append(cell if cell else "")
                    rows.append(cells)
                result.append(rows)
            return result
    except Exception:
        pass
    
    return []


def extract_all_text_grid(page):
    """
    Extract all text from a page and arrange into a grid structure.
    Groups text by approximate Y position into rows,
    then by X position into columns.
    """
    blocks = page.get_text("dict", sort=True)["blocks"]
    
    # Collect all text spans with positions
    text_items = []
    for block in blocks:
        if block["type"] != 0:
            continue
        for line in block.get("lines", []):
            spans = line.get("spans", [])
            for span in spans:
                text = span.get("text", "").strip()
                if text:
                    bbox = span.get("bbox", [0, 0, 0, 0])
                    text_items.append({
                        "text": text,
                        "x": bbox[0],
                        "y": bbox[1],
                        "x2": bbox[2],
                        "y2": bbox[3],
                        "size": span.get("size", 11),
                        "font": span.get("font", ""),
                        "flags": span.get("flags", 0),
                    })
    
    if not text_items:
        return []
    
    # Group by Y position (within tolerance)
    text_items.sort(key=lambda t: (t["y"], t["x"]))
    
    y_tolerance = 3  # Points
    rows = []
    current_row = [text_items[0]]
    current_y = text_items[0]["y"]
    
    for item in text_items[1:]:
        if abs(item["y"] - current_y) <= y_tolerance:
            current_row.append(item)
        else:
            rows.append(sorted(current_row, key=lambda t: t["x"]))
            current_row = [item]
            current_y = item["y"]
    
    if current_row:
        rows.append(sorted(current_row, key=lambda t: t["x"]))
    
    # Convert to simple text grid
    result_rows = []
    for row in rows:
        # Merge items that are close together horizontally
        merged = []
        current_text = row[0]["text"]
        current_x2 = row[0]["x2"]
        
        for item in row[1:]:
            gap = item["x"] - current_x2
            if gap < 15:  # Close together, merge
                current_text += " " + item["text"]
                current_x2 = item["x2"]
            else:
                merged.append(current_text)
                current_text = item["text"]
                current_x2 = item["x2"]
        
        merged.append(current_text)
        result_rows.append(merged)
    
    return result_rows


def normalize_column_count(rows):
    """Ensure all rows have the same number of columns."""
    if not rows:
        return rows
    max_cols = max(len(row) for row in rows)
    for row in rows:
        while len(row) < max_cols:
            row.append("")
    return rows


def auto_detect_number(value):
    """Try to convert a string value to a number for Excel."""
    if not isinstance(value, str):
        return value
    
    text = value.strip()
    if not text:
        return ""
    
    # Remove currency symbols and thousands separators for number detection
    cleaned = re.sub(r'[$€£¥,]', '', text)
    cleaned = cleaned.strip()
    
    try:
        if '.' in cleaned:
            return float(cleaned)
        return int(cleaned)
    except (ValueError, TypeError):
        pass
    
    # Percentage
    if cleaned.endswith('%'):
        try:
            return float(cleaned[:-1]) / 100
        except (ValueError, TypeError):
            pass
    
    return text


def write_rows_to_sheet(ws, rows, merge_cells=True, start_row=1):
    """Write row data to a worksheet with formatting."""
    if not rows:
        return start_row
    
    rows = normalize_column_count(rows)
    
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    header_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    header_font = Font(bold=True, size=10)
    normal_font = Font(size=10)
    
    for row_idx, row in enumerate(rows):
        excel_row = start_row + row_idx
        for col_idx, cell_value in enumerate(row):
            cell = ws.cell(row=excel_row, column=col_idx + 1)
            cell.value = auto_detect_number(cell_value) if cell_value else ""
            cell.border = thin_border
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            
            # First row gets header styling
            if row_idx == 0:
                cell.fill = header_fill
                cell.font = header_font
            else:
                cell.font = normal_font
    
    # Auto-size columns
    for col_idx in range(len(rows[0]) if rows else 0):
        max_len = 0
        col_letter = get_column_letter(col_idx + 1)
        for row in rows:
            if col_idx < len(row):
                cell_len = len(str(row[col_idx] or ""))
                max_len = max(max_len, cell_len)
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 8), 50)
    
    # Handle merge_cells for identical adjacent cells
    if merge_cells and len(rows) > 1:
        for col_idx in range(len(rows[0])):
            merge_start = start_row
            for row_idx in range(1, len(rows)):
                excel_row = start_row + row_idx
                current_val = rows[row_idx][col_idx] if col_idx < len(rows[row_idx]) else ""
                prev_val = rows[row_idx - 1][col_idx] if col_idx < len(rows[row_idx - 1]) else ""
                
                if current_val != prev_val or current_val == "" or row_idx == len(rows) - 1:
                    # End of a potential merge region
                    if row_idx == len(rows) - 1 and current_val == prev_val and current_val != "":
                        merge_end = excel_row
                    else:
                        merge_end = excel_row - 1
                    
                    if merge_end > merge_start:
                        try:
                            ws.merge_cells(
                                start_row=merge_start,
                                start_column=col_idx + 1,
                                end_row=merge_end,
                                end_column=col_idx + 1
                            )
                        except Exception:
                            pass
                    
                    merge_start = excel_row
    
    return start_row + len(rows)


def main():
    parser = argparse.ArgumentParser(description="Convert PDF to Excel (.xlsx)")
    parser.add_argument("input_pdf", help="Path to input PDF file")
    parser.add_argument("output_xlsx", help="Path to output XLSX file")
    parser.add_argument("--mode", choices=["tables", "all"], default="tables",
                        help="Extraction mode: tables (detect tables) or all (all text)")
    parser.add_argument("--merge-cells", dest="merge_cells", action="store_true", default=True,
                        help="Detect and merge identical adjacent cells")
    parser.add_argument("--no-merge-cells", dest="merge_cells", action="store_false",
                        help="Do not merge cells")
    parser.add_argument("--sheet-per-page", dest="sheet_per_page", action="store_true", default=True,
                        help="Create a separate sheet for each page")
    parser.add_argument("--single-sheet", dest="sheet_per_page", action="store_false",
                        help="Put all data on a single sheet")
    parser.add_argument("--json", action="store_true", default=False,
                        help="Output result as JSON")
    
    args = parser.parse_args()
    
    if not os.path.exists(args.input_pdf):
        result = {"success": False, "error": f"Input file not found: {args.input_pdf}"}
        if args.json:
            print(json.dumps(result))
        else:
            print(f"Error: {result['error']}", file=sys.stderr)
        sys.exit(1)
    
    try:
        pdf_doc = fitz.open(args.input_pdf)
        wb = Workbook()
        
        # Remove default sheet if we're creating per-page sheets
        if args.sheet_per_page:
            wb.remove(wb.active)
        
        total_pages = len(pdf_doc)
        tables_found = 0
        current_row = 1  # For single-sheet mode
        
        for page_idx in range(total_pages):
            page = pdf_doc[page_idx]
            page_label = f"Page {page_idx + 1}"
            
            if args.mode == "tables":
                page_tables = find_tables_on_page(page)
                
                if not page_tables:
                    # Fallback: extract all text as grid
                    text_rows = extract_all_text_grid(page)
                    if text_rows:
                        page_tables = [text_rows]
                
                for table_idx, table_rows in enumerate(page_tables):
                    if not table_rows:
                        continue
                    
                    tables_found += 1
                    
                    if args.sheet_per_page:
                        sheet_name = page_label
                        if table_idx > 0:
                            sheet_name = f"Page {page_idx + 1} T{table_idx + 1}"
                        # Truncate sheet name to 31 chars (Excel limit)
                        sheet_name = sheet_name[:31]
                        ws = wb.create_sheet(title=sheet_name)
                        write_rows_to_sheet(ws, table_rows, args.merge_cells)
                    else:
                        if not wb.sheetnames:
                            ws = wb.create_sheet(title="Data")
                        else:
                            ws = wb.active
                        
                        # Add page header
                        if current_row > 1:
                            current_row += 1  # Empty row separator
                        ws.cell(row=current_row, column=1, value=page_label)
                        ws.cell(row=current_row, column=1).font = Font(bold=True, size=11)
                        current_row += 1
                        
                        current_row = write_rows_to_sheet(ws, table_rows, args.merge_cells, current_row)
            
            else:  # mode == "all"
                text_rows = extract_all_text_grid(page)
                
                if text_rows:
                    if args.sheet_per_page:
                        ws = wb.create_sheet(title=page_label[:31])
                        write_rows_to_sheet(ws, text_rows, args.merge_cells)
                    else:
                        if not wb.sheetnames:
                            ws = wb.create_sheet(title="Data")
                        else:
                            ws = wb.active
                        
                        if current_row > 1:
                            current_row += 1
                        ws.cell(row=current_row, column=1, value=page_label)
                        ws.cell(row=current_row, column=1).font = Font(bold=True, size=11)
                        current_row += 1
                        
                        current_row = write_rows_to_sheet(ws, text_rows, args.merge_cells, current_row)
        
        # Ensure at least one sheet exists
        if not wb.sheetnames:
            ws = wb.create_sheet(title="Sheet1")
            ws.cell(row=1, column=1, value="No extractable content found in this PDF.")
        
        wb.save(args.output_xlsx)
        pdf_doc.close()
        
        file_size = os.path.getsize(args.output_xlsx)
        
        result = {
            "success": True,
            "output": args.output_xlsx,
            "pages": total_pages,
            "tables_found": tables_found,
            "sheets": len(wb.sheetnames) if hasattr(wb, 'sheetnames') else 0,
            "file_size": file_size,
            "mode": args.mode,
        }
        
        if args.json:
            print(json.dumps(result))
        else:
            print(f"Converted {total_pages} pages to {args.output_xlsx} ({file_size} bytes, {tables_found} tables)")
    
    except Exception as e:
        result = {"success": False, "error": str(e), "traceback": traceback.format_exc()}
        if args.json:
            print(json.dumps(result))
        else:
            print(f"Error: {e}", file=sys.stderr)
            traceback.print_exc()
        sys.exit(1)


if __name__ == "__main__":
    main()
